"""Exporta os dados do TFT (CommunityDragon) para planilhas com colunas legíveis.

Gera saida/planilhas/tft-set<N>-dados.xlsx (uma aba por tabela, mais LEIA-ME e
Dicionário) e um CSV por aba em saida/planilhas/csv/.

Uso:
  python3 -m venv materiais/aula-tft/.cache/venv
  materiais/aula-tft/.cache/venv/bin/pip install openpyxl
  materiais/aula-tft/.cache/venv/bin/python materiais/aula-tft/exportar-planilhas.py
Sem openpyxl, gera só os CSVs.
"""

import csv
import datetime
import json
import os
import re
import subprocess

AQUI = os.path.dirname(os.path.abspath(__file__))
CACHE = os.path.join(AQUI, ".cache")
SAIDA = os.path.join(AQUI, "saida", "planilhas")
SET = os.environ.get("TFT_SET", "18")
BASE = "https://raw.communitydragon.org/latest"
os.makedirs(os.path.join(SAIDA, "csv"), exist_ok=True)


def baixar(url, arquivo, obrigatorio=True):
    destino = os.path.join(CACHE, arquivo)
    if not os.path.exists(destino):
        # O CDN recusa o user-agent padrão do Python; curl é aceito.
        r = subprocess.run(["curl", "-s", "-f", "-o", destino, url])
        if r.returncode != 0:
            if obrigatorio:
                raise SystemExit(f"Falha ao baixar {url}")
            return None
    with open(destino, encoding="utf-8") as f:
        return json.load(f)


dados = baixar(f"{BASE}/cdragon/tft/pt_br.json", "pt_br.json")
mapa = baixar(f"{BASE}/game/data/maps/shipping/map22/map22.bin.json", "map22.json")
versao = baixar(f"{BASE}/content-metadata.json", "content-metadata.json")["version"].split("+")[0]
hoje = datetime.date.today().isoformat()

conf = next(s for s in dados["setData"] if s.get("mutator") == f"TFTSet{SET}")
cfg_set = mapa[f"Maps/Shipping/Map22/Sets/TFTSet{SET}"]
por_api = {i["apiName"]: i for i in dados["items"]}

# ---------- tradução de nomes internos ----------
TRADUCAO = {
    "AD": "Dano de ataque",
    "AP": "Poder de habilidade",
    "AS": "Velocidade de ataque (%)",
    "AttackSpeed": "Velocidade de ataque",
    "Armor": "Armadura",
    "MagicResist": "Resistência mágica",
    "Health": "Vida",
    "BonusPercentHP": "Vida máxima extra (%)",
    "PercentMaxHP": "Vida máxima extra (%)",
    "CritChance": "Chance de crítico (%)",
    "CritDamage": "Dano crítico",
    "Mana": "Mana",
    "ManaRegen": "Mana por segundo",
    "CombatStartMana": "Mana no início do combate",
    "FlatManaRestore": "Mana extra por ataque",
    "Omnivamp": "Omnivampirismo",
    "StatOmnivamp": "Omnivampirismo",
    "LifeSteal": "Roubo de vida",
    "DamageAmp": "Amplificação de dano",
    "BonusDamage": "Dano bônus",
    "Durability": "Durabilidade (redução de dano)",
    "DamageReduction": "Redução de dano",
    "HexRange": "Alcance em hexes",
    "Duration": "Duração (s)",
    "ShieldDuration": "Duração do escudo (s)",
    "ShieldHealthPercent": "Escudo (% da vida)",
    "HealthThreshold": "Limite de vida (%)",
    "ICD": "Tempo de recarga interno (s)",
    "BurnPercent": "Queimadura (% da vida por segundo)",
    "GrievousWoundsPercent": "Ferimento (redução de cura, %)",
    "MRShred": "Fragmentação (redução de resistência mágica, %)",
    "ArmorReductionPercent": "Fatiamento (redução de armadura, %)",
    "ARReductionAmount": "Fatiamento (redução de armadura, %)",
    "AttackSpeedPerStack": "Velocidade de ataque por acúmulo (%)",
    "MaxStacks": "Acúmulos máximos",
    "StackCap": "Acúmulos máximos",
    "InstantGold": "Ouro imediato",
    "Gold": "Ouro",
    "XP": "Experiência",
    "MinUnits": "Mínimo de unidades",
}


def traduz(chave):
    if chave.startswith("{"):
        return "(nome interno codificado)"
    return TRADUCAO.get(chave, "")


def limpa(texto, valores=None, sequencia=None):
    """Remove marcações do jogo e troca @Variavel@ pelo valor quando ele existe."""
    if not texto:
        return ""
    valores = valores or {}
    fila = list(sequencia or [])

    def troca(m):
        chave, vezes100 = m.group(1), m.group(2)
        if chave == "MinUnits" and fila:
            return str(fila.pop(0))
        v = valores.get(chave)
        if isinstance(v, (int, float)):
            v = v * 100 if vezes100 else v
            return f"{v:g}".replace(".", ",")
        return "▢"

    t = texto.replace("\\r", " ").replace("\\n", " ").replace("\r", " ").replace("\n", " ")
    t = re.sub(r"<br\s*/?>", " ", t)
    t = re.sub(r"<row>", " • ", t)
    t = re.sub(r"%i:[^%]+%", "", t)
    t = re.sub(r"\{\{[^}]+\}\}", "", t)
    t = re.sub(r"@([^@*]+)(\*100)?@", troca, t)
    t = re.sub(r"<[^>]+>", "", t)
    return re.sub(r"\s+", " ", t).strip()


def num(v, casas=4):
    if isinstance(v, float):
        return round(v, casas)
    return v


# ---------- campeões ----------
def raiz_bin(api):
    b = baixar(f"{BASE}/game/characters/{api.lower()}.cdtb.bin.json", f"{api.lower()}.json", obrigatorio=False)
    if not b:
        return {}
    return next((v for k, v in b.items() if k.endswith("CharacterRecords/Root")), {})


def base(r, chave):
    v = r.get(chave)
    return v.get("baseValue") if isinstance(v, dict) else None


MULT_VIDA = {1: 1, 2: 1.8, 3: 3.24}
MULT_DANO = {1: 1, 2: 1.5, 3: 2.25}


def por_estrela(s, loja):
    vida, dano, vel = s.get("hp"), s.get("damage"), s.get("attackSpeed")
    if not loja or vida is None or dano is None or vel is None:
        return [None] * 7
    dps = lambda e: num(dano * MULT_DANO[e] * vel, 1)
    return [dps(1), num(vida * MULT_VIDA[2], 0), num(dano * MULT_DANO[2], 1), dps(2),
            num(vida * MULT_VIDA[3], 0), num(dano * MULT_DANO[3], 1), dps(3)]


campeoes, camp_carac, hab_var = [], [], []
for c in sorted(conf["champions"], key=lambda c: (c["cost"], c["name"])):
    s, r = c["stats"], raiz_bin(c["apiName"])
    loja = bool(c["traits"]) and 1 <= c["cost"] <= 5
    campeoes.append([
        c["name"], c["apiName"], "Campeão da loja" if loja else "Invocação ou objeto", c["cost"],
        ", ".join(c["traits"]), num(s.get("hp")), num(s.get("armor")), num(s.get("magicResist")),
        num(s.get("damage")), num(s.get("attackSpeed"), 3), *por_estrela(s, loja), num(s.get("critChance")), num(s.get("critMultiplier"), 3),
        num(s.get("initialMana")), num(s.get("mana")), num(s.get("range")),
        num(base(r, "attackRangeModifiable")), num(base(r, "baseMoveSpeedModifiable")), num(base(r, "attackSpeedRatioModifiable"), 3),
        c["ability"]["name"], limpa(c["ability"]["desc"]),
    ])
    for t in c["traits"]:
        camp_carac.append([c["name"], c["cost"], t])
    for v in c["ability"].get("variables") or []:
        vals = v.get("value") or []
        hab_var.append([c["name"], c["ability"]["name"], v.get("name"), *[num(vals[i]) if i < len(vals) else None for i in (1, 2, 3)]])

# ---------- características ----------
ESTILO = {1: "Bronze", 3: "Prata", 4: "Único", 5: "Ouro", 6: "Prismático"}
membros = {}
for c in conf["champions"]:
    for t in c["traits"]:
        membros.setdefault(t, []).append(c["name"])
caracteristicas, carac_niveis = [], []
for t in sorted(conf["traits"], key=lambda t: t["name"]):
    efs = t.get("effects") or []
    limiares = [e["minUnits"] for e in efs]
    caracteristicas.append([
        t["name"], t["apiName"], len(efs), " / ".join(map(str, limiares)),
        ", ".join(sorted(set(membros.get(t["name"], [])))), len(set(membros.get(t["name"], []))),
        limpa(t.get("desc"), (efs[0].get("variables") if efs else {}), limiares),
    ])
    for e in efs:
        variaveis = e.get("variables") or {}
        if not variaveis:
            carac_niveis.append([t["name"], e["minUnits"], min(e["maxUnits"], 99), ESTILO.get(e.get("style"), e.get("style")), None, None, None])
        for k, v in variaveis.items():
            carac_niveis.append([t["name"], e["minUnits"], min(e["maxUnits"], 99), ESTILO.get(e.get("style"), e.get("style")), k, traduz(k), num(v)])

# ---------- itens e aumentos ----------
NOME_LISTA = {
    "Common_Items": "Itens comuns",
    f"Set{SET}_Items": f"Itens do set {SET} (emblemas e outros)",
    f"TFTSet{SET}_Items_Charms": f"Amuletos do set {SET}",
    "Common_Items_Augments": "Aumentos comuns",
    "Common_Items_Augments_Set17": "Aumentos comuns (lista do set 17)",
    f"Set{SET}_Items_Augments": f"Aumentos do set {SET}",
}
origem = {}
for caminho in cfg_set["itemLists"]:
    lista = mapa.get(caminho) or {}
    nome = NOME_LISTA.get(caminho.split("/")[-1], caminho.split("/")[-1])
    for item in lista.get("mItems", []):
        origem.setdefault(item.split("/")[-1], nome)

componentes = {x for i in dados["items"] if i.get("composition") for x in i["composition"]}


def tipo_item(i):
    a, desc = i["apiName"], i.get("desc") or ""
    if i.get("composition"):
        return "Completo"
    if a in componentes:
        return "Componente"
    if "Emblem" in a:
        return "Emblema"
    if "Artifact" in a or "Ornn" in a:
        return "Artefato"
    if "Radiant" in a:
        return "Radiante"
    if "Item de Suporte" in desc:
        return "Suporte"
    if "Consumable" in a:
        return "Consumível"
    if "Charm" in a or "Amuleto" in origem.get(a, ""):
        return "Amuleto"
    return "Outro"


def efeito_como_texto(i, k, v):
    return v * 100 if isinstance(v, (int, float)) and re.search("@" + re.escape(k) + r"\*100@", i.get("desc") or "") else v


STATS = [("AD", "Dano de ataque"), ("AP", "Poder de habilidade"), ("AS", "Velocidade de ataque (%)"), ("Armor", "Armadura"),
         ("MagicResist", "Resistência mágica"), ("Health", "Vida"), ("CritChance", "Chance de crítico (%)"), ("ManaRegen", "Mana por segundo")]
def nome_legivel(i):
    nome = i["name"]
    return f"(sem tradução) {i['apiName']}" if re.search(r"_name$", nome, re.I) else nome


def equivalente_padrao(api, ef):
    if ef or not api.startswith("DA_"):
        return ""
    padrao = por_api.get("TFT_Item_" + api[3:])
    return padrao["name"] if padrao and padrao.get("effects") else ""


itens, itens_ef, aumentos, aum_ef = [], [], [], []
vistos = set()
for api in conf["items"]:
    i = por_api.get(api)
    if not i or "Debug" in api or not i.get("name") or api in vistos:
        continue
    vistos.add(api)
    ef = i.get("effects") or {}
    receita = " + ".join(por_api[x]["name"] for x in (i.get("composition") or []) if x in por_api)
    itens.append([nome_legivel(i), api, tipo_item(i), origem.get(api, ""), receita, "sim" if i.get("unique") else "não",
                  ", ".join(i.get("associatedTraits") or []), "sim" if ef else "não", equivalente_padrao(api, ef),
                  *[num(ef.get(k)) for k, _ in STATS], limpa(i.get("desc"), ef)])
    for k, v in ef.items():
        if v is not None:
            itens_ef.append([nome_legivel(i), api, k, traduz(k), num(v), num(efeito_como_texto(i, k, v))])


def tier_aumento(icone):
    m = re.search(r"[_-](i{1,3})(?:[._])", (icone or "").lower())
    return {"i": "Prata", "ii": "Ouro", "iii": "Prismático"}.get(m.group(1)) if m else ""


for api in conf["augments"]:
    a = por_api.get(api)
    if not a or not a.get("name"):
        continue
    ef = a.get("effects") or {}
    aumentos.append([nome_legivel(a), api, tier_aumento(a.get("icon")), origem.get(api, ""),
                     ", ".join(a.get("associatedTraits") or []), limpa(a.get("desc"), ef)])
    for k, v in ef.items():
        if v is not None:
            aum_ef.append([nome_legivel(a), api, k, traduz(k), num(v), num(efeito_como_texto(a, k, v))])

# ---------- loja, rodadas, dano ----------
chances = mapa[cfg_set["DropRateTables"]["Shop"]]["mDropRatesByLevel"]
loja = []
for nivel, linha in enumerate(chances, start=1):
    p = next(v for v in linha.values() if isinstance(v, list))
    loja.append([nivel, *[round(x * 100, 1) for x in p], round(5 * p[2], 2), round(5 * p[3], 2), round(5 * p[4], 2)])

NOME_RODADA = {"Combat": "Combate contra jogador", "Carousel": "Carrossel", "Krugs": "Monstros (Krugues)", "Wolves": "Monstros (Lobos)",
               "Raptors": "Monstros (Aves)", "Intro_Carousel": "Carrossel inicial", "Intro1": "Introdução (monstros)",
               "Intro2": "Introdução (monstros)", "Intro3": "Introdução (monstros)"}
rodadas = []
for e, est in enumerate(mapa[cfg_set["StageRoundData"]]["stages"], start=1):
    for r, caminho in enumerate(est["mRounds"], start=1):
        nome = caminho.split("/")[-1]
        info = mapa.get(caminho) or {}
        curto = nome.replace("TFT_Round_", "")
        rodadas.append([f"{e}-{r}", e, r, NOME_RODADA.get(curto, curto), nome,
                        (info.get("mPlanning") or {}).get("mDuration"), (info.get("mCombat") or {}).get("mDuration")])

tabela_dano = next(o["mConstants"] for o in mapa.values() if isinstance(o, dict) and (o.get("mConstants") or {}).get("PlayerDamageStage3", {}).get("mValue") == 6
                   and o["mConstants"].get("PlayerDamageStage2", {}).get("mValue") == 2)
dano = [[e, (tabela_dano.get(f"PlayerDamageStage{e}") or {}).get("mValue", 0), 1] for e in range(1, 8)]

# ---------- tabelas e dicionário ----------
fonte = f"CommunityDragon (dados extraídos dos arquivos do jogo; não oficial). Set {SET}, versão {versao}, baixado em {hoje}."
TABELAS = [
    ("Campeoes", "Todos os campeões e invocações do set: atributos de 1 estrela e, para campeões da loja, vida, dano e DPS calculados para 2 e 3 estrelas.", [
        ("Nome", "Nome em português."), ("Nome interno", "Identificador nos arquivos do jogo."),
        ("Tipo", "Campeão comprável na loja ou invocação/objeto (plantas, monstros, bigornas…)."), ("Custo", "Custo em ouro (valores acima de 5 são de objetos especiais)."),
        ("Características", "Origens e classes do campeão."), ("Vida", "Vida com 1 estrela."), ("Armadura", "Reduz dano físico: dano × 100/(100 + armadura) (regra conhecida; conferir)."),
        ("Resistência mágica", "Reduz dano mágico pela mesma fórmula da armadura."), ("Dano de ataque", "Dano por ataque básico com 1 estrela."),
        ("Velocidade de ataque", "Ataques por segundo."),
        ("DPS 1 estrela", "Dano de ataque × velocidade de ataque, sem crítico, armadura nem habilidade."),
        ("Vida 2 estrelas", "Vida × 1,8 (Wiki do LoL, TFT Codex, TFT Ninja; não oficial: conferir no jogo)."),
        ("Dano de ataque 2 estrelas", "Dano de ataque × 1,5 (Wiki do LoL, TFT Ninja; conferir no jogo)."),
        ("DPS 2 estrelas", "Dano de ataque 2 estrelas × velocidade de ataque (a velocidade não muda com estrelas)."),
        ("Vida 3 estrelas", "Vida × 3,24 (= 1,8²)."), ("Dano de ataque 3 estrelas", "Dano de ataque × 2,25 (= 1,5²)."),
        ("DPS 3 estrelas", "Dano de ataque 3 estrelas × velocidade de ataque."), ("Chance de crítico", "Probabilidade de crítico (0,25 = 25%)."),
        ("Multiplicador de crítico", "Dano do crítico em relação ao normal (1,4 = 140%)."), ("Mana inicial", "Mana no começo do combate."),
        ("Mana máxima", "Mana para conjurar a habilidade."), ("Alcance (hexes)", "Alcance mostrado no jogo."),
        ("Alcance (unidades do jogo)", "Mesmo alcance em unidades internas."), ("Velocidade de movimento (unidades/s)", "Unidades internas por segundo."),
        ("Razão de velocidade de ataque", "Fator interno que converte bônus de velocidade de ataque (uso exato não documentado)."),
        ("Habilidade", "Nome da habilidade."), ("Descrição da habilidade", "Texto do jogo; ▢ = número que não vem nos arquivos públicos."),
    ], campeoes),
    ("Campeoes_Caracteristicas", "Uma linha por par campeão–característica (bom para tabela dinâmica).", [
        ("Campeão", "Nome do campeão."), ("Custo", "Custo em ouro."), ("Característica", "Origem ou classe."),
    ], camp_carac),
    ("Habilidades_Variaveis", "Números das habilidades quando os arquivos públicos trazem (poucos campeões).", [
        ("Campeão", "Nome do campeão."), ("Habilidade", "Nome da habilidade."), ("Variável", "Nome interno do número."),
        ("Valor 1 estrela", "Valor com 1 estrela."), ("Valor 2 estrelas", "Valor com 2 estrelas."), ("Valor 3 estrelas", "Valor com 3 estrelas."),
    ], hab_var),
    ("Caracteristicas", "Origens e classes do set.", [
        ("Nome", "Nome em português."), ("Nome interno", "Identificador nos arquivos."), ("Número de níveis", "Quantos degraus de ativação."),
        ("Limiares", "Quantas unidades ativam cada nível."), ("Campeões", "Quem tem a característica."), ("Quantidade de campeões", "Total de campeões com ela."),
        ("Descrição", "Texto do jogo com os limiares preenchidos; ▢ = número não publicado."),
    ], caracteristicas),
    ("Caracteristicas_Niveis", "Uma linha por nível e por número de cada característica.", [
        ("Característica", "Nome."), ("Mínimo de unidades", "Unidades para ativar o nível."), ("Máximo de unidades", "Até quantas unidades vale o nível (99 = sem limite)."),
        ("Cor do nível", "Bronze, prata, ouro, prismático (mapeamento estimado a partir do código de estilo)."),
        ("Efeito (nome interno)", "Nome do número nos arquivos; entre chaves = nome codificado."), ("Efeito (tradução)", "Tradução quando conhecida."),
        ("Valor", "Número cru dos arquivos (pode ser fração: 0,25 = 25%)."),
    ], carac_niveis),
    ("Itens", f"Itens ativos no set {SET} (sem itens de teste).", [
        ("Nome", "Nome em português."), ("Nome interno", "Identificador nos arquivos."),
        ("Tipo", "Componente, Completo, Emblema, Artefato, Radiante, Suporte, Consumível, Amuleto ou Outro (pela receita, lista e nome)."),
        ("Lista de origem", "Lista do jogo que inclui o item no set."), ("Receita", "Componentes que formam o item."), ("Único", "Só um por campeão."),
        ("Características associadas", "Características ligadas ao item (emblemas)."),
        ("Tem números próprios", "“não” = os arquivos públicos não trazem os números desta versão do item."),
        ("Item padrão equivalente", "Para itens do set sem números: item padrão de mesmo nome que tem números (procure-o nesta aba). Qual versão o jogo usa: conferir."),
        *[(nome, f"Atributo “{nome}” do item (vazio = não tem). Dano de ataque vem como fração: 0,1 = 10%.") for _, nome in STATS],
        ("Descrição", "Texto do jogo com os números preenchidos; ▢ = número não publicado."),
    ], itens),
    ("Itens_Efeitos", "Todos os números de cada item, um por linha (filtre pelo item ou pelo efeito).", [
        ("Item", "Nome."), ("Nome interno do item", "Identificador."), ("Efeito (nome interno)", "Nome do número nos arquivos."),
        ("Efeito (tradução)", "Tradução quando conhecida."), ("Valor cru", "Como está nos arquivos."),
        ("Valor como aparece no texto", "Multiplicado por 100 quando o texto do jogo mostra em porcentagem."),
    ], itens_ef),
    ("Aumentos", f"Aumentos (augments) ativos no set {SET}.", [
        ("Nome", "Nome em português."), ("Nome interno", "Identificador."), ("Tier", "Prata, Ouro ou Prismático (estimado pelo nome do ícone; vazio = não identificado)."),
        ("Lista de origem", "Lista do jogo que inclui o aumento."), ("Características associadas", "Características ligadas ao aumento."),
        ("Descrição", "Texto do jogo com os números preenchidos; ▢ = número não publicado."),
    ], aumentos),
    ("Aumentos_Efeitos", "Todos os números de cada aumento, um por linha.", [
        ("Aumento", "Nome."), ("Nome interno do aumento", "Identificador."), ("Efeito (nome interno)", "Nome do número."),
        ("Efeito (tradução)", "Tradução quando conhecida."), ("Valor cru", "Como está nos arquivos."), ("Valor como aparece no texto", "Multiplicado por 100 quando o texto mostra em %."),
    ], aum_ef),
    ("Loja_Chances", "Chance de cada custo aparecer em cada vaga da loja, por nível do jogador.", [
        ("Nível", "Nível do jogador."), *[(f"Custo {k} (%)", f"Chance de uma vaga da loja ser de custo {k}.") for k in range(1, 6)],
        ("Média de custo 3 por loja", "5 vagas × chance."), ("Média de custo 4 por loja", "5 vagas × chance."), ("Média de custo 5 por loja", "5 vagas × chance."),
    ], loja),
    ("Rodadas", "Sequência de rodadas de cada estágio.", [
        ("Rodada", "Estágio-rodada, como no jogo (2-1)."), ("Estágio", "Número do estágio."), ("Posição", "Posição da rodada no estágio."),
        ("Tipo", "Combate contra jogador, carrossel ou monstros."), ("Nome interno", "Identificador da rodada."),
        ("Planejamento (s)", "Duração da fase de compra e posicionamento."), ("Combate (s)", "Duração da fase de combate."),
    ], rodadas),
    ("Dano_por_Estagio", "Dano ao jogador por derrota = dano do estágio + 1 por unidade inimiga viva (regra oficial da Riot desde a atualização 14.8).", [
        ("Estágio", "Número do estágio."),
        ("Dano fixo por derrota", "Tabela encontrada nos arquivos (igual à do tft-boost, atualização 18.1). OP.GG e TFT Ninja dão 5 e 8 nos estágios 3 e 4: conferir."),
        ("Dano por inimigo vivo", "Somar este valor por unidade inimiga viva, sem importar custo ou estrelas."),
    ], dano),
]

leia_me = [
    ["Dados do TFT para análise"],
    [fonte],
    [],
    ["Como usar no Google Sheets: Arquivo → Importar → Fazer upload → este .xlsx. Cada aba vira uma página."],
    ["Abas com sufixo _Efeitos ou _Niveis estão em formato longo (uma linha por número): use filtro ou tabela dinâmica."],
    ["▢ nas descrições = número que o jogo calcula mas os arquivos públicos não trazem (leia no tooltip do jogo)."],
    ["Regras de economia (juros, bônus de sequência, XP) não estão nos arquivos e por isso não aparecem aqui."],
    [],
    ["Aba", "Conteúdo", "Linhas"],
    *[[nome, desc, len(linhas)] for nome, desc, _, linhas in TABELAS],
]
dicionario = [[nome, col, sig] for nome, _, cols, _ in TABELAS for col, sig in cols]

# ---------- escrita ----------
for nome, _, cols, linhas in TABELAS:
    with open(os.path.join(SAIDA, "csv", f"{nome}.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow([c for c, _ in cols])
        w.writerows(linhas)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("openpyxl não instalado: gerei só os CSVs.")
else:
    wb = Workbook()
    ws = wb.active
    ws.title = "LEIA-ME"
    for linha in leia_me:
        ws.append(linha)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 90

    def aba(titulo, cabecalho, linhas):
        ws = wb.create_sheet(titulo)
        ws.append(cabecalho)
        for linha in linhas:
            ws.append(linha)
        for celula in ws[1]:
            celula.font = Font(bold=True, color="FFFFFF")
            celula.fill = PatternFill("solid", fgColor="7C3AED")
            celula.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        for i, titulo_col in enumerate(cabecalho, start=1):
            maior = max([len(str(titulo_col))] + [len(str(l[i - 1])) for l in linhas[:200] if i - 1 < len(l) and l[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, maior + 2), 60)

    aba("Dicionario", ["Aba", "Coluna", "Significado"], dicionario)
    for nome, _, cols, linhas in TABELAS:
        aba(nome, [c for c, _ in cols], linhas)
    arquivo = os.path.join(SAIDA, f"tft-set{SET}-dados.xlsx")
    wb.save(arquivo)
    print(f"Planilha: {arquivo}")

for nome, _, _, linhas in TABELAS:
    print(f"{nome}: {len(linhas)} linhas")
